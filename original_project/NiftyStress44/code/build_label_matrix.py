"""
Builds risk_label_matrix.xlsx: an 11 sectors x 4 stress events grid showing
the risk label at a glance, color-coded (red=High, amber=Medium, green=Low).

Input : ../outputs/labeled_risk_dataset.csv
Output: ../outputs/risk_label_matrix.xlsx
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FONT = "Arial"
INPUT_FILE = "../outputs/labeled_risk_dataset.csv"
OUTPUT_FILE = "../outputs/risk_label_matrix.xlsx"

EVENT_ORDER = ["COVID-19 Crash", "Russia-Ukraine Shock", "RBI Rate Hike Cycle", "US Fed Tightening"]
LABEL_FILLS = {
    "High": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
    "Medium": PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid"),
    "Low": PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid"),
}

df = pd.read_csv(INPUT_FILE)

pivot = df.pivot(index="Sector", columns="Event_Name", values="Risk_Label")
pivot = pivot[[c for c in EVENT_ORDER if c in pivot.columns]]

wb = Workbook()
ws = wb.active
ws.title = "Risk Label Matrix"

ws.cell(row=1, column=1, value="Sector")
for j, event in enumerate(pivot.columns, start=2):
    ws.cell(row=1, column=j, value=event)

for i, (sector, row) in enumerate(pivot.iterrows(), start=2):
    ws.cell(row=i, column=1, value=sector).font = Font(name=FONT, bold=True)
    for j, event in enumerate(pivot.columns, start=2):
        label = row[event]
        cell = ws.cell(row=i, column=j, value=label)
        cell.font = Font(name=FONT)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = LABEL_FILLS.get(label, PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"))

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
for c in range(1, len(pivot.columns) + 2):
    cell = ws.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions[get_column_letter(c)].width = 22
ws.freeze_panes = "B2"

# Notes sheet -- flags if any label is entirely absent (e.g. Low never occurring)
ws2 = wb.create_sheet("Notes")
ws2["A1"] = "Note"
ws2["A1"].font = Font(name=FONT, bold=True)
present_labels = set(df["Risk_Label"].unique())
missing_labels = {"High", "Medium", "Low"} - present_labels
if missing_labels:
    note = (f"The following risk label(s) do not occur in this dataset: {sorted(missing_labels)}. "
             "This can be a legitimate finding (e.g. all analyzed stress windows were severe enough "
             "that no sector stayed within the 'Low' band) rather than a data error -- verify against "
             "your MDD threshold rule and stress window definitions before treating it as a bug.")
else:
    note = "All three risk labels (High, Medium, Low) are present in this dataset."
ws2["A2"] = note
ws2["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws2.column_dimensions["A"].width = 100

wb.save(OUTPUT_FILE)
print(f"Saved {OUTPUT_FILE}")
print(pivot)
if missing_labels:
    print(f"\nNOTE: label(s) {sorted(missing_labels)} do not appear in the dataset. See Notes sheet.")
